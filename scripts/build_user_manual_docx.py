#!/usr/bin/env python3
from __future__ import annotations

import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
LOCAL_DEPS = ROOT / ".pydeps"
if LOCAL_DEPS.exists():
    sys.path.insert(0, str(LOCAL_DEPS))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SOFTWARE_NAME = "Wamo Chat"
SOFTWARE_VERSION = "V1.0"
DOCUMENT_TITLE = f"{SOFTWARE_NAME} 使用说明书"
OUTPUT_DIR = ROOT / "output" / "doc"
OUTPUT_DOCX = OUTPUT_DIR / "software-user-manual-wamo-chat-v1.0-android-zh-hans.docx"
SCREENSHOT_DIR = Path(os.environ.get("WAMO_SCREENSHOT_DIR", "/Users/soulc/Downloads/android-screen"))

FEATURES = [
    {
        "title": "1、安装确认页",
        "image": "安装确认页",
        "flow": "Android 文件管理器 -> 点击 Wamo Chat 安装包 -> 安装确认页",
        "description": "用户在 Android 设备上点击 Wamo Chat 安装包后，系统展示安装确认界面，用于确认应用来源并执行安装操作。",
        "usage": "在该页面点击“安装”后，系统开始部署应用；安装完成后点击“打开”，即可进入软件首次启动页面。",
    },
    {
        "title": "2、首次启动页",
        "image": "首次启动页",
        "flow": "安装完成 -> 点击“打开” -> 首次启动页 -> 进入主界面",
        "description": "首次启动页用于完成应用初次加载，展示软件品牌标识和启动状态，帮助用户确认软件已经成功安装。",
        "usage": "用户等待启动流程完成后，软件自动跳转到主界面；若首次进入需要读取基础配置，也会在该阶段完成初始化。",
    },
    {
        "title": "3、主界面完整图",
        "image": "主界面完整图",
        "flow": "首次启动页 -> 主界面",
        "description": "主界面是软件的核心操作页面，集中提供会话列表、消息显示区域、输入框以及底部功能入口。",
        "usage": "用户可在该页面直接输入问题、创建新会话、打开设置、上传文件或图片，并从底部区域切换模型和功能模块。",
    },
    {
        "title": "4、通用设置页",
        "image": "通用设置页",
        "flow": "主界面 -> 打开设置 -> 通用设置页",
        "description": "通用设置页用于调整语言、主题、网络代理等基础参数，帮助用户根据设备环境优化软件显示和连接体验。",
        "usage": "用户修改对应参数后返回主界面即可生效；该页面通常由主界面进入，是后续模型配置和功能使用的基础设置入口。",
    },
    {
        "title": "5、模型提供方配置页",
        "image": "模型提供方配置页",
        "flow": "主界面 -> 设置 -> 模型提供方配置页",
        "description": "模型提供方配置页用于填写 API Host、API Key 等连接信息，是软件调用大模型服务的前置页面。",
        "usage": "用户选择目标提供方并保存参数后，软件即可与对应模型服务建立连接；保存成功后可继续设置默认模型或直接返回主界面。",
    },
    {
        "title": "6、默认模型设置页",
        "image": "默认模型设置页",
        "flow": "设置页面 -> 默认模型设置页",
        "description": "默认模型设置页用于指定新建会话时优先使用的模型，便于统一后续问答和生成任务的默认能力。",
        "usage": "用户选择合适的对话模型并保存后，新建会话会自动调用该模型；如需切换模型，可再次进入该页面重新设置。",
    },
    {
        "title": "7、新会话发送与回复页",
        "image": "新会话发送与回复页",
        "flow": "主界面 -> 新建会话 -> 输入问题 -> 发送 -> 返回回复",
        "description": "新会话发送与回复页展示了软件最核心的智能对话功能，用户可在此输入文本问题并查看模型返回结果。",
        "usage": "用户在输入框中录入问题后点击发送按钮，系统将请求发送给当前模型并返回答案；在同一页面中还可继续进行多轮追问。",
    },
    {
        "title": "8、文件上传与问答页",
        "image": "文件上传与问答页",
        "flow": "会话页面 -> 添加文件 -> 选择文档 -> 输入问题 -> 获取回答",
        "description": "文件上传与问答页用于上传文本类资料，并围绕上传内容进行定向提问，帮助用户快速完成文档理解和摘要提取。",
        "usage": "用户先在会话页面选择文件，等待系统完成解析后再提交问题；软件将结合文件内容返回相应答案，并保持与当前会话关联。",
    },
    {
        "title": "9、图片上传与视觉问答页",
        "image": "图片上传与视觉问答页",
        "flow": "会话页面 -> 添加图片 -> 输入视觉问题 -> 返回识别结果",
        "description": "图片上传与视觉问答页用于对图片进行识别、描述和内容理解，是软件的多模态问答功能页面。",
        "usage": "用户上传图片并输入问题后，系统调用支持视觉能力的模型分析图片内容，再以文本形式返回识别结果或解释说明。",
    },
    {
        "title": "10、联网问答页",
        "image": "联网问答页",
        "flow": "会话页面 -> 启用联网问答 -> 输入问题 -> 查看结果",
        "description": "联网问答页用于处理具有时效性的查询任务，软件会结合在线检索结果生成回答，提升实时信息获取能力。",
        "usage": "用户开启联网问答后输入问题并发送，系统先执行检索再整理结果返回；完成后可继续在同一会话中追问相关内容。",
    },
    {
        "title": "11、图片生成页",
        "image": "图片生成页",
        "flow": "主界面 -> 打开图片生成器 -> 输入提示词 -> 生成图片",
        "description": "图片生成页用于根据文字提示生成新图像，适合进行创意草图、海报素材和视觉概念的快速生成。",
        "usage": "用户输入提示词并提交后，系统调用图像生成模型创建结果图片；生成完成后可直接在当前页面查看并继续调整提示词。",
    },
    {
        "title": "12、导出与分享页",
        "image": "导出:分享页",
        "flow": "会话页面 -> 打开导出或分享入口 -> 系统分享面板",
        "description": "导出与分享页用于将当前会话内容输出到其他应用或本地文件，方便保存记录、继续编辑或发送给他人。",
        "usage": "用户在会话页面选择导出或分享操作后，系统弹出分享面板；完成目标应用选择后，即可将聊天内容继续传递到外部环境。",
    },
    {
        "title": "13、文档解析设置页",
        "image": "文档解析设置页",
        "flow": "设置页面 -> 文档解析设置页 -> 选择解析方式 -> 保存",
        "description": "文档解析设置页用于配置文件处理方式，帮助用户在不同网络和解析能力条件下选择合适的文档理解方案。",
        "usage": "用户在该页面切换解析方式并保存后，后续上传文件时将按照新的配置执行解析；设置完成后返回会话页面即可继续测试。",
    },
]


def set_run_font(run, *, font_name: str = "宋体", size_pt: float = 11, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)


def add_field(paragraph, field_code: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "1"

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run = paragraph.add_run()
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)
    set_run_font(run, font_name="宋体", size_pt=9)


def add_bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "808080")
    p_bdr.append(bottom)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.4)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = False

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    add_header(section)


def add_header(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(15.2), WD_TAB_ALIGNMENT.RIGHT)

    left = paragraph.add_run(f"{SOFTWARE_NAME} {SOFTWARE_VERSION}")
    set_run_font(left, font_name="宋体", size_pt=9)
    paragraph.add_run("\t")
    add_field(paragraph, "PAGE")
    add_bottom_border(paragraph)


def add_title(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(20)
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run(text)
    set_run_font(run, font_name="黑体", size_pt=18, bold=True)


def add_section_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, font_name="黑体", size_pt=13, bold=True)


def add_body_paragraph(doc: Document, text: str, *, first_line_indent: bool = True, left_indent_cm: float = 0.0, space_after_pt: float = 4.0) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(space_after_pt)
    if first_line_indent:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    if left_indent_cm:
        paragraph.paragraph_format.left_indent = Cm(left_indent_cm)
    run = paragraph.add_run(text)
    set_run_font(run, font_name="宋体", size_pt=11)


def add_image(doc: Document, image_path: Path) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run.add_picture(str(image_path), height=Cm(14.0))


def normalize_name(name: str) -> str:
    text = Path(name).stem
    for token in ["Android-", "Android_", "（", "）", "(", ")", "`", "'", '"', "：", ":", "-", "_", " "]:
        text = text.replace(token, "")
    return text.lower()


def build_screenshot_index() -> dict[str, Path]:
    index: dict[str, Path] = {}
    if not SCREENSHOT_DIR.exists():
        return index
    for path in SCREENSHOT_DIR.iterdir():
        if not path.is_file():
            continue
        if path.suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp"}:
            continue
        index[normalize_name(path.name)] = path
    return index


def resolve_image(image_key: str, index: dict[str, Path]) -> Path | None:
    key = normalize_name(image_key)
    if key in index:
        return index[key]
    for filename_key, path in index.items():
        if key in filename_key or filename_key in key:
            return path
    return None


def add_intro_page(doc: Document) -> None:
    add_title(doc, DOCUMENT_TITLE)
    add_section_heading(doc, "一、软件介绍")
    add_body_paragraph(
        doc,
        "Wamo Chat 是一款运行于 Android 平台的移动端智能对话软件，支持 AI 模型配置、多轮会话、文件问答、图片理解、联网检索、图片生成与内容分享等主要能力。",
    )
    add_body_paragraph(
        doc,
        "本软件面向移动办公与日常智能助手场景，帮助用户在手机端完成信息检索、内容创作、资料处理与图像生成，提高移动使用场景下的工作效率。",
    )
    add_section_heading(doc, "二、主要功能介绍")
    add_body_paragraph(
        doc,
        "以下按照 Android 端实际操作顺序，对安装启动、主界面浏览、参数设置、智能问答、文件与图片处理、联网问答、图片生成以及导出分享等主要功能逐项进行介绍。",
    )


def add_feature_page(doc: Document, feature: dict[str, str]) -> None:
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(4)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.keep_with_next = True
    run = heading.add_run(feature["title"])
    set_run_font(run, font_name="黑体", size_pt=12, bold=True)

    for text in [
        f"页面跳转：{feature['flow']}。",
        f"功能说明：{feature['description']}",
        f"使用说明：{feature['usage']}",
        "+完整清晰的界面图",
    ]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(text)
        set_run_font(run, font_name="宋体", size_pt=11)


def build_manual() -> tuple[Path, list[str]]:
    screenshot_index = build_screenshot_index()
    missing: list[str] = []

    doc = Document()
    configure_document(doc)
    add_intro_page(doc)

    for index, feature in enumerate(FEATURES):
        doc.add_page_break()
        add_feature_page(doc, feature)
        image_path = resolve_image(feature["image"], screenshot_index)
        if image_path is None:
            missing.append(feature["image"])
            add_body_paragraph(doc, f"未匹配到截图：{feature['image']}", first_line_indent=False)
            continue
        add_image(doc, image_path)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    return OUTPUT_DOCX, missing


def main() -> None:
    output_path, missing = build_manual()
    print(f"DOCX generated: {output_path}")
    if missing:
        print("Missing screenshots:")
        for item in missing:
            print(f"  - {item}")
    else:
        print("Missing screenshots: none")


if __name__ == "__main__":
    main()
