#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


INCLUDE_EXTENSIONS = {".ts", ".tsx", ".js", ".jsx", ".cjs", ".mjs"}
EXCLUDE_TOKENS = (".test.", ".spec.", "__tests__", ".bk")
LINES_PER_PAGE = 50
MAX_PAGES = 60
HALF_PAGES = 30
LINE_PREVIEW_WIDTH = 80
DOC_LINE_MAX_CHARS = 92


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate source code DOCX document.")
    parser.add_argument(
        "--repo-root",
        type=Path,
        default=Path(__file__).resolve().parents[1],
        help="Repository root path",
    )
    parser.add_argument(
        "--source-dir",
        type=Path,
        default=Path("src"),
        help="Relative source directory under repo root",
    )
    parser.add_argument(
        "--software-name",
        type=str,
        default="软件名称",
        help="Software name",
    )
    parser.add_argument(
        "--software-version",
        type=str,
        default="V1.0",
        help="Software version label",
    )
    parser.add_argument(
        "--output-path",
        type=Path,
        default=None,
        help="Output DOCX path",
    )
    return parser.parse_args()


def load_software_info(repo_root: Path) -> tuple[str, str]:
    package_file = repo_root / "package.json"
    if not package_file.exists():
        return "软件名称", "V1.0"

    try:
        data = json.loads(package_file.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return "软件名称", "V1.0"

    name = data.get("productName") or data.get("name") or "软件名称"
    version = data.get("version") or "1.0"
    return name, f"V{version}"


def iter_source_files(source_root: Path) -> Iterable[Path]:
    paths: list[Path] = []
    for path in source_root.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix not in INCLUDE_EXTENSIONS:
            continue
        text_path = str(path)
        if any(token in text_path for token in EXCLUDE_TOKENS):
            continue
        paths.append(path)
    paths.sort()
    return paths


def normalize_line(text: str) -> str:
    line = text.rstrip("\n\r")
    line = line.replace("\t", "    ")
    return line[:LINE_PREVIEW_WIDTH]


def fit_doc_width(text: str) -> str:
    if len(text) <= DOC_LINE_MAX_CHARS:
        return text
    return f"{text[: DOC_LINE_MAX_CHARS - 3]}..."


def collect_code_lines(repo_root: Path, source_root: Path) -> list[str]:
    lines: list[str] = []
    for file_path in iter_source_files(source_root):
        rel = file_path.relative_to(repo_root).as_posix()
        lines.append(fit_doc_width(f"// ===== FILE: {rel} ====="))
        content = file_path.read_text(encoding="utf-8", errors="replace").splitlines()
        for idx, raw_line in enumerate(content, start=1):
            lines.append(fit_doc_width(f"{idx:04d} | {normalize_line(raw_line)}"))
        lines.append(fit_doc_width(f"// ===== END FILE: {rel} ====="))
    return lines


def split_lines_for_submission(all_lines: list[str]) -> list[str]:
    required_lines = MAX_PAGES * LINES_PER_PAGE
    if len(all_lines) > required_lines:
        front = all_lines[: HALF_PAGES * LINES_PER_PAGE]
        back = all_lines[-HALF_PAGES * LINES_PER_PAGE :]
        return front + back

    selected = list(all_lines)
    remainder = len(selected) % LINES_PER_PAGE
    if remainder != 0:
        selected.extend("//" for _ in range(LINES_PER_PAGE - remainder))
    return selected


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def build_document(header_title: str, selected_lines: list[str], output_path: Path) -> tuple[int, int]:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Courier New"
    normal_style.font.size = Pt(8)

    header = section.header.paragraphs[0]
    header.text = header_title
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        run.font.name = "Songti SC"
        run.font.size = Pt(10)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    add_page_field(footer)
    footer.add_run(" 页")
    for run in footer.runs:
        run.font.name = "Songti SC"
        run.font.size = Pt(10)

    page_count = len(selected_lines) // LINES_PER_PAGE

    for idx, line in enumerate(selected_lines, start=1):
        paragraph = doc.add_paragraph(line)
        pf = paragraph.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        if idx % LINES_PER_PAGE == 0 and idx != len(selected_lines):
            doc.add_page_break()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return page_count, len(selected_lines)


def main() -> None:
    args = parse_args()
    repo_root = args.repo_root.resolve()
    source_root = (repo_root / args.source_dir).resolve()

    inferred_name, inferred_version = load_software_info(repo_root)
    software_name = args.software_name if args.software_name != "软件名称" else inferred_name
    software_version = args.software_version if args.software_version != "V1.0" else inferred_version
    header_title = f"{software_name} {software_version}"

    output_path = args.output_path
    if output_path is None:
        today = date.today().isoformat()
        output_path = repo_root / "output" / "doc" / f"源代码文档-{today}.docx"
    else:
        output_path = output_path.resolve()

    all_lines = collect_code_lines(repo_root=repo_root, source_root=source_root)
    if not all_lines:
        raise SystemExit(f"未在 {source_root} 中找到可用源码文件")

    selected_lines = split_lines_for_submission(all_lines)
    page_count, final_line_count = build_document(
        header_title=header_title,
        selected_lines=selected_lines,
        output_path=output_path,
    )

    print(f"Output: {output_path}")
    print(f"Header: {header_title}")
    print(f"Source lines collected: {len(all_lines)}")
    print(f"Source lines exported: {final_line_count}")
    print(f"Pages exported: {page_count}")
    if len(all_lines) > MAX_PAGES * LINES_PER_PAGE:
        print("Mode: 前30页+后30页")
    else:
        print("Mode: 全量源码（不足60页）")


if __name__ == "__main__":
    main()
